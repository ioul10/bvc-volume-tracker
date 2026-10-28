"""
Indicateur de liquidité par valeur, construit à partir des bulletins.

Quatre dimensions, calculées sur les N dernières séances (médianes) :
  - Volume de fond   : volume journalier hors grosses transactions (blocs)
  - Fréquence        : part des séances où la valeur a été traitée
  - Écart achat/vente: (meilleure offre - meilleure demande) / milieu, à la clôture
  - Impact (Amihud)  : |variation %| / volume en millions de MAD
Le score (0-100) combine les rangs en percentile de chaque mesure.
"""
import io

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from bvc_excel import BLOC_PAR_CONTRAT_MIN, BLOC_VOLUME_MIN

POIDS = {"volume": 0.30, "frequence": 0.20, "ecart": 0.20, "impact": 0.30}


def _classe(score):
    if score >= 75:
        return "Très liquide"
    if score >= 50:
        return "Liquide"
    if score >= 25:
        return "Peu liquide"
    return "Illiquide"


def liquidite(h, vj, nb_seances=20):
    """Retourne (tableau par valeur, série marché par séance)."""
    dates = sorted(vj["Date"].unique())[-nb_seances:]
    n = len(dates)

    # Volume de fond = volume total - volume des blocs
    phy = h["Physionomie"]
    phy = phy[phy["Date"].isin(dates)].copy()
    phy["bloc"] = (phy["Volume_MAD"] >= BLOC_VOLUME_MIN) & (phy["Volume_MAD"] / phy["Contrats"] >= BLOC_PAR_CONTRAT_MIN)
    fond = (phy[~phy["bloc"]].groupby(["Date", "Valeur"])["Volume_MAD"].sum()
            .rename("Volume_fond").reset_index())
    blocs = phy[phy["bloc"]].groupby(["Date", "Valeur"])["Volume_MAD"].sum().rename("Volume_blocs").reset_index()

    d = vj[vj["Date"].isin(dates)].merge(fond, on=["Date", "Valeur"], how="left").merge(blocs, on=["Date", "Valeur"], how="left")
    d[["Volume_fond", "Volume_blocs"]] = d[["Volume_fond", "Volume_blocs"]].fillna(0)
    cot = h["Cotation"][["Date", "Valeur", "Demande", "Offre"]]
    d = d.merge(cot, on=["Date", "Valeur"], how="left")
    mid = (d["Offre"] + d["Demande"]) / 2
    d["Ecart_pct"] = np.where((d["Offre"] > 0) & (d["Demande"] > 0), (d["Offre"] - d["Demande"]) / mid * 100, np.nan)
    d["Amihud"] = np.where(d["Volume_MAD"] > 0, d["Var_pct"].abs() / (d["Volume_MAD"] / 1e6), np.nan)
    d["Amplitude_pct"] = np.where(d["Volume_MAD"] > 0, (d["Plus_haut"] - d["Plus_bas"]) / d["Cloture"] * 100, np.nan)

    g = d.groupby("Valeur").agg(
        Secteur=("Secteur_nom", "last"),
        Volume_fond_median=("Volume_fond", "median"),
        Volume_moyen=("Volume_MAD", "mean"),
        Part_blocs_pct=("Volume_blocs", "sum"),
        Volume_total=("Volume_MAD", "sum"),
        Seances_traitees=("Volume_MAD", lambda s: int((s > 0).sum())),
        Contrats_moyens=("Contrats", "mean"),
        Ecart_median_pct=("Ecart_pct", "median"),
        Amihud_median=("Amihud", "median"),
        Amplitude_median_pct=("Amplitude_pct", "median"),
        Cloture=("Cloture", "last"),
    ).reset_index()
    g["Part_blocs_pct"] = np.where(g["Volume_total"] > 0, g["Part_blocs_pct"] / g["Volume_total"] * 100, 0)
    g["Frequence_pct"] = g["Seances_traitees"] / n * 100
    g["MAD_pour_1pct_impact"] = np.where(g["Amihud_median"] > 0, 1e6 / g["Amihud_median"], np.nan)

    # Rangs en percentile (100 = le plus liquide). Valeur manquante = la pire note.
    # method="max" : des valeurs à égalité reçoivent la meilleure note (ex. 100 % de régularité = 20/20)
    r_vol = g["Volume_fond_median"].rank(pct=True, method="max")
    r_freq = g["Frequence_pct"].rank(pct=True, method="max")
    r_ecart = (-g["Ecart_median_pct"]).rank(pct=True, method="max").fillna(0)
    r_impact = (-g["Amihud_median"]).rank(pct=True, method="max").fillna(0)
    # Points de chaque ingrédient (le total fait le score sur 100)
    g["Pts_volume"] = 100 * POIDS["volume"] * r_vol
    g["Pts_frequence"] = 100 * POIDS["frequence"] * r_freq
    g["Pts_ecart"] = 100 * POIDS["ecart"] * r_ecart
    g["Pts_impact"] = 100 * POIDS["impact"] * r_impact
    g["Score"] = g[["Pts_volume", "Pts_frequence", "Pts_ecart", "Pts_impact"]].sum(axis=1)
    g["Classe"] = g["Score"].apply(_classe)
    g = g.drop(columns=["Volume_total"]).sort_values("Score", ascending=False).reset_index(drop=True)
    g.insert(0, "Rang", range(1, len(g) + 1))

    # Liquidité du marché séance par séance
    marche = d.groupby("Date").agg(
        Volume_total=("Volume_MAD", "sum"),
        Volume_fond=("Volume_fond", "sum"),
        Volume_blocs=("Volume_blocs", "sum"),
        Valeurs_traitees=("Volume_MAD", lambda s: int((s > 0).sum())),
        Valeurs_cotees=("Valeur", "count"),
        Ecart_median_pct=("Ecart_pct", "median"),
    ).reset_index()
    marche["Largeur_pct"] = marche["Valeurs_traitees"] / marche["Valeurs_cotees"] * 100
    return g, marche


def volume_fond_jour(h, dates=None):
    """Volume hors grosses transactions, par valeur et par séance."""
    phy = h["Physionomie"]
    if dates is not None:
        phy = phy[phy["Date"].isin(dates)]
    bloc = (phy["Volume_MAD"] >= BLOC_VOLUME_MIN) & (phy["Volume_MAD"] / phy["Contrats"] >= BLOC_PAR_CONTRAT_MIN)
    return phy[~bloc].groupby(["Date", "Valeur"])["Volume_MAD"].sum().rename("Volume_fond").reset_index()


def ajouter_feuille(xlsx_bytes, df, nom="Liquidite"):
    """Ajoute le tableau de liquidité au classeur Excel produit par bvc_excel."""
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    if nom in wb.sheetnames:
        del wb[nom]
    ws = wb.create_sheet(nom, 1)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in row])
    fill = PatternFill("solid", fgColor="1F3A5F")
    formats = {"Volume": "#,##0", "MAD_pour": "#,##0", "_pct": '0.0" %"', "Amihud": "0.000",
               "Score": "0", "Contrats": "#,##0", "Cloture": "#,##0.00"}
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=1, column=j)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), fill
        ws.column_dimensions[get_column_letter(j)].width = 26 if col in ("Valeur", "Secteur") else 16
        fmt = next((f for k, f in formats.items() if k in col), None)
        if fmt:
            for i in range(2, ws.max_row + 1):
                ws.cell(row=i, column=j).number_format = fmt
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
