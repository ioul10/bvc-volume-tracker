"""
Construit l'historique (plusieurs bulletins + ancien Excel), calcule les
indicateurs de volume et écrit un classeur Excel mis en forme.
"""
import io

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAW_SHEETS = ["Marche", "Cotation", "Physionomie", "Indices"]

SECTEURS = {
    "AGR": "Agroalimentaire", "ASS": "Assurances", "BAN": "Banques", "BMC": "BTP & matériaux",
    "BOI": "Boissons", "CHI": "Chimie", "DIS": "Distributeurs", "ELC": "Électricité",
    "IAG": "Industrie agricole", "IBE": "Ingénierie", "IMM": "Immobilier", "LH": "Loisirs & hôtels",
    "LSI": "Informatique", "MIN": "Mines", "PEG": "Pétrole & gaz", "PHA": "Pharmacie",
    "SAN": "Santé", "SDT": "Services de transport", "SFA": "Sociétés de financement",
    "SP": "Papier", "SPH": "Holdings", "SPI": "Foncières (SPI)", "TCO": "Télécoms", "TRP": "Transport",
}

# Seuils (modifiables)
SEUIL_INSTITUTIONNEL = 250_000     # MAD par contrat
SEUIL_PARTICULIERS = 60_000        # MAD par contrat
SEUIL_VARIATION = 2.0              # % sur la période
SEUIL_TENDANCE = 0.20              # +/-20 % de volume
BLOC_VOLUME_MIN = 5_000_000        # MAD sur un niveau de prix
BLOC_PAR_CONTRAT_MIN = 1_000_000   # MAD par contrat


# ---------------------------------------------------------------- historique
def bulletins_to_frames(parsed):
    """Liste de résultats parse_bulletin -> dict de DataFrames bruts."""
    marche, cot, phy, idx = [], [], [], []
    for r in parsed:
        d = pd.Timestamp(r["date"])
        masi = r["indices"][r["indices"]["Indice"] == "MASI"]
        marche.append({
            "Date": d,
            "Volume_MAD": r["totaux"]["volume"],
            "Quantite": r["totaux"]["quantite"],
            "Contrats": r["totaux"]["contrats"],
            "MASI": masi["Valeur"].iloc[0] if len(masi) else None,
            "MASI_var_pct": masi["Var_jour_pct"].iloc[0] if len(masi) else None,
        })
        cot.append(r["cotation"])
        phy.append(r["physionomie"])
        idx.append(r["indices"])
    return {
        "Marche": pd.DataFrame(marche),
        "Cotation": pd.concat(cot, ignore_index=True) if cot else pd.DataFrame(),
        "Physionomie": pd.concat(phy, ignore_index=True) if phy else pd.DataFrame(),
        "Indices": pd.concat(idx, ignore_index=True) if idx else pd.DataFrame(),
    }


def read_existing(xlsx_file):
    """Relit les feuilles brutes d'un classeur produit par cet outil."""
    out = {}
    xls = pd.ExcelFile(xlsx_file)
    for s in RAW_SHEETS:
        if s in xls.sheet_names:
            df = pd.read_excel(xls, s)
            df["Date"] = pd.to_datetime(df["Date"])
            out[s] = df
    return out


def merge_history(old, new):
    """Fusionne : pour une date présente dans les deux, la nouvelle version gagne."""
    merged = {}
    for s in RAW_SHEETS:
        a, b = old.get(s, pd.DataFrame()), new.get(s, pd.DataFrame())
        if a.empty:
            merged[s] = b.copy()
        elif b.empty:
            merged[s] = a.copy()
        else:
            a = a[~a["Date"].isin(b["Date"].unique())]
            merged[s] = pd.concat([a, b], ignore_index=True)
        if not merged[s].empty:
            merged[s] = merged[s].sort_values("Date").reset_index(drop=True)
    return merged


# ---------------------------------------------------------------- indicateurs
def volumes_journaliers(h):
    """Une ligne par valeur et par séance avec volume, VWAP, ticket, ratio 20 séances."""
    phy, cot = h["Physionomie"], h["Cotation"]
    agg = (phy.groupby(["Date", "Valeur"], as_index=False)
              .agg(Volume_MAD=("Volume_MAD", "sum"), Quantite=("Quantite", "sum"), Contrats=("Contrats", "sum")))
    df = cot.merge(agg, on=["Date", "Valeur"], how="left")
    for c in ["Volume_MAD", "Quantite", "Contrats"]:
        df[c] = df[c].fillna(0)
    df["VWAP"] = np.where(df["Quantite"] > 0, df["Volume_MAD"] / df["Quantite"].replace(0, np.nan), np.nan)
    df["Ticket_moyen"] = np.where(df["Contrats"] > 0, df["Volume_MAD"] / df["Contrats"].replace(0, np.nan), np.nan)
    df["Cloture_vs_VWAP_pct"] = (df["Cloture"] / df["VWAP"] - 1) * 100
    tot = df.groupby("Date")["Volume_MAD"].transform("sum")
    df["Part_marche_pct"] = df["Volume_MAD"] / tot * 100
    df["Secteur_nom"] = df["Secteur"].map(SECTEURS).fillna(df["Secteur"])

    # Moyenne des 20 séances PRÉCÉDENTES (jours sans échange = 0)
    df = df.sort_values(["Valeur", "Date"])
    df["Moy_20_seances"] = (df.groupby("Valeur")["Volume_MAD"]
                              .transform(lambda s: s.shift(1).rolling(20, min_periods=1).mean()))
    df["Ratio_vs_moy20"] = df["Volume_MAD"] / df["Moy_20_seances"].replace(0, np.nan)
    df["Nb_seances_historique"] = df.groupby("Valeur").cumcount()
    cols = ["Date", "Valeur", "Secteur_nom", "Statut", "Cours_ref", "Ouverture", "Cloture", "Plus_haut", "Plus_bas",
            "Var_pct", "Volume_MAD", "Quantite", "Contrats", "VWAP", "Cloture_vs_VWAP_pct", "Ticket_moyen",
            "Part_marche_pct", "Moy_20_seances", "Ratio_vs_moy20", "Nb_seances_historique",
            "Plus_haut_annuel", "Plus_bas_annuel"]
    return df[cols].sort_values(["Date", "Volume_MAD"], ascending=[True, False]).reset_index(drop=True)


def _flux(ticket):
    if pd.isna(ticket):
        return ""
    if ticket >= SEUIL_INSTITUTIONNEL:
        return "Institutionnel"
    if ticket <= SEUIL_PARTICULIERS:
        return "Particuliers"
    return "Mixte"


def _signal(var, tendance, volume):
    up, dn = var >= SEUIL_VARIATION, var <= -SEUIL_VARIATION
    rising = tendance is None or np.isnan(tendance) or tendance >= 1 + SEUIL_TENDANCE
    falling = tendance is not None and not np.isnan(tendance) and tendance <= 1 - SEUIL_TENDANCE
    if up and rising:
        return "Hausse confirmée"
    if up and falling:
        return "Hausse sans volume"
    if up:
        return "Hausse, volume stable"
    if dn and rising:
        return "Baisse sous volume"
    if dn and falling:
        return "Baisse sans volume"
    if dn:
        return "Baisse, volume stable"
    if abs(var) < 1 and volume >= 20_000_000:
        return "Absorption"
    return "Neutre"


def signaux(vj, nb_seances=5):
    """Synthèse par valeur sur les N dernières séances."""
    dates = sorted(vj["Date"].unique())[-nb_seances:]
    w = vj[vj["Date"].isin(dates)]
    n = len(dates)
    k = n - n // 2  # première moitié (ex. 3 séances sur 5)
    rows = []
    for val, g in w.groupby("Valeur"):
        g = g.set_index("Date").reindex(dates)
        v = g["Volume_MAD"].fillna(0).values
        tot = v.sum()
        if tot <= 0:
            continue
        contrats = g["Contrats"].fillna(0).sum()
        ref, close = g["Cours_ref"].iloc[0], g["Cloture"].iloc[-1]
        var = (close / ref - 1) * 100 if ref and close else np.nan
        early, late = v[:k].mean(), v[k:].mean() if n > 1 else np.nan
        tendance = late / early if early > 0 else np.nan
        med = np.median(v)
        row = {
            "Valeur": val,
            "Secteur": g["Secteur_nom"].dropna().iloc[-1] if g["Secteur_nom"].notna().any() else "",
            "Volume_periode_MAD": tot,
            "Part_periode_pct": 0.0,
            "Var_periode_pct": var,
            "Tendance_volume": tendance,
            "Pic_vs_mediane": (v.max() / med) if med > 0 else np.nan,
            "Ratio_dernier_jour_vs_moy20": g["Ratio_vs_moy20"].iloc[-1],
            "Ticket_moyen": tot / contrats if contrats else np.nan,
            "Flux": "",
            "Cloture": close,
            "VWAP_periode": tot / g["Quantite"].fillna(0).sum(),
            "Ecart_plus_haut_annuel_pct": (close / g["Plus_haut_annuel"].iloc[-1] - 1) * 100 if g["Plus_haut_annuel"].iloc[-1] else np.nan,
        }
        for d, x in zip(dates, v):
            row[pd.Timestamp(d).strftime("%d/%m")] = x
        row["Flux"] = _flux(row["Ticket_moyen"])
        row["Signal"] = _signal(var if not np.isnan(var) else 0, tendance, tot)
        rows.append(row)
    s = pd.DataFrame(rows)
    if not s.empty:
        s["Part_periode_pct"] = s["Volume_periode_MAD"] / s["Volume_periode_MAD"].sum() * 100
        s = s.sort_values("Volume_periode_MAD", ascending=False).reset_index(drop=True)
    return s


def grosses_transactions(h, vj):
    phy = h["Physionomie"].copy()
    phy["MAD_par_contrat"] = phy["Volume_MAD"] / phy["Contrats"]
    b = phy[(phy["Volume_MAD"] >= BLOC_VOLUME_MIN) & (phy["MAD_par_contrat"] >= BLOC_PAR_CONTRAT_MIN)]
    b = b.merge(vj[["Date", "Valeur", "Cloture", "Volume_MAD"]].rename(columns={"Volume_MAD": "Volume_jour"}),
                on=["Date", "Valeur"], how="left")
    b["Part_du_jour_pct"] = b["Volume_MAD"] / b["Volume_jour"] * 100
    b["Prix_vs_cloture_pct"] = (b["Prix"] / b["Cloture"] - 1) * 100
    b["Position"] = np.select(
        [b["Prix_vs_cloture_pct"].abs() < 0.05, b["Prix_vs_cloture_pct"] > 0],
        ["À la clôture", "Au-dessus de la clôture"], "En dessous de la clôture")
    cols = ["Date", "Valeur", "Prix", "Quantite", "Volume_MAD", "Contrats", "MAD_par_contrat",
            "Part_du_jour_pct", "Cloture", "Prix_vs_cloture_pct", "Position"]
    return b[cols].sort_values(["Date", "Volume_MAD"], ascending=[False, False]).reset_index(drop=True)


def matrice_volumes(vj):
    """Valeurs en lignes, séances en colonnes (millions de MAD)."""
    m = vj.pivot_table(index="Valeur", columns="Date", values="Volume_MAD", aggfunc="sum", fill_value=0) / 1e6
    m.columns = [pd.Timestamp(c).strftime("%d/%m/%Y") for c in m.columns]
    m["Total"] = m.sum(axis=1)
    return m.sort_values("Total", ascending=False).reset_index()


# ---------------------------------------------------------------- Excel
FORMATS = {
    "Date": "dd/mm/yyyy",
    "pct": '0.00" %"',
    "mad": "#,##0",
    "prix": "#,##0.00",
    "ratio": '0.0"×"',
    "mmad": "#,##0.0",
}


def _fmt_for(col):
    c = col.lower()
    if col == "Date":
        return FORMATS["Date"]
    if c.endswith("_pct"):
        return FORMATS["pct"]
    if c.startswith(("tendance", "pic", "ratio")):
        return FORMATS["ratio"]
    if any(k in c for k in ("volume", "ticket", "moy_20", "mad_par", "quantite", "contrats")):
        return FORMATS["mad"]
    if any(k in c for k in ("cours", "ouverture", "cloture", "plus_", "prix", "vwap", "demande", "offre", "masi", "valeur_")):
        return FORMATS["prix"]
    return None


def build_workbook(h, nb_seances=5):
    """Retourne (bytes du fichier Excel, dict des tableaux calculés)."""
    vj = volumes_journaliers(h)
    sig = signaux(vj, nb_seances)
    blocs = grosses_transactions(h, vj)
    mat = matrice_volumes(vj)

    sheets = [
        ("Signaux", sig),
        ("Matrice_volumes_MMAD", mat),
        ("Volumes_jour", vj),
        ("Grosses_transactions", blocs),
        ("Marche", h["Marche"]),
        ("Cotation", h["Cotation"]),
        ("Physionomie", h["Physionomie"]),
        ("Indices", h["Indices"]),
    ]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets:
            df.to_excel(xw, sheet_name=name, index=False)
    buf.seek(0)

    wb = load_workbook(buf)
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    for name, df in sheets:
        ws = wb[name]
        ws.freeze_panes = "C2" if name in ("Signaux", "Volumes_jour", "Matrice_volumes_MMAD") else "B2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for j, col in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=j)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            fmt = FORMATS["mmad"] if name == "Matrice_volumes_MMAD" and col != "Valeur" else _fmt_for(str(col))
            if name == "Signaux" and "/" in str(col):
                fmt = FORMATS["mad"]
            if fmt:
                for i in range(2, ws.max_row + 1):
                    ws.cell(row=i, column=j).number_format = fmt
            width = max(10, min(28, len(str(col)) + 2))
            if col in ("Valeur", "Signal", "Secteur", "Secteur_nom", "Position", "Indice"):
                width = 26
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.row_dimensions[1].height = 32

    # Carte de chaleur sur les volumes
    def heat(ws, first_col, last_col):
        if ws.max_row > 1 and last_col >= first_col:
            rng = f"{get_column_letter(first_col)}2:{get_column_letter(last_col)}{ws.max_row}"
            ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                              end_type="percentile", end_value=95, end_color="2E86C1"))

    heat(wb["Matrice_volumes_MMAD"], 2, len(mat.columns) - 1)
    day_cols = [j for j, c in enumerate(sig.columns, start=1) if "/" in str(c)]
    if day_cols:
        heat(wb["Signaux"], day_cols[0], day_cols[-1])

    # Couleur des signaux
    colors = {"Hausse confirmée": "C6EFCE", "Baisse sous volume": "FFC7CE",
              "Hausse sans volume": "FFEB9C", "Baisse sans volume": "FFEB9C", "Absorption": "DDEBF7"}
    if "Signal" in sig.columns:
        j = list(sig.columns).index("Signal") + 1
        ws = wb["Signaux"]
        for i in range(2, ws.max_row + 1):
            c = ws.cell(row=i, column=j)
            if c.value in colors:
                c.fill = PatternFill("solid", fgColor=colors[c.value])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), {"volumes_jour": vj, "signaux": sig, "blocs": blocs, "matrice": mat}
